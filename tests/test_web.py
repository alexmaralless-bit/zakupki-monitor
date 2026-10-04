from __future__ import annotations

import json
import warnings
from collections.abc import Iterator
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any, cast

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from zakupki_monitor.ai import DemoProvider  # noqa: F401  (провайдер demo доступен по имени)
from zakupki_monitor.appsettings import load_runtime
from zakupki_monitor.collector import evaluate
from zakupki_monitor.config import Settings
from zakupki_monitor.models import CrmExport, Profile, Rule, Run, Tender
from zakupki_monitor.profiles import ProfileSpec, save_profile
from zakupki_monitor.scoring import RuleSpec
from zakupki_monitor.search_url import SearchParams
from zakupki_monitor.store import upsert_tender
from zakupki_monitor.web.app import create_app, deadline_info

from .conftest import make_tender_entry

warnings.filterwarnings("ignore", category=DeprecationWarning)
TODAY = date.today()


def state(client: TestClient) -> Any:
    return cast(Any, client.app).state


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    settings = Settings(db_path=tmp_path / "web.db", ai_provider="yandex")
    app = create_app(settings, start_scheduler=False)
    with app.state.factory() as session:
        profile = save_profile(
            session,
            ProfileSpec(
                name="Газ",
                search=SearchParams(search_string="газопровод"),
                rules=[RuleSpec(type="keyword", value="газопровод", weight=2)],
            ),
        )
        rows = [
            make_tender_entry("32600000001", initial_price=Decimal("5000000"), deadline_at=TODAY + timedelta(days=2)),
            make_tender_entry("32600000002", object="Закупка канцтоваров", law=44, initial_price=Decimal("90000"),
                              deadline_at=TODAY + timedelta(days=15)),
            make_tender_entry("32600000003", object="Проектирование газопровода низкого давления",
                              initial_price=Decimal("1000000"), deadline_at=TODAY + timedelta(days=9)),
        ]  # fmt: skip
        for entry in rows:
            tender = upsert_tender(session, entry).tender
            evaluate(session, tender, profile, TODAY)
        session.add(Run(profile_id=profile.id, pages=2, found=30, new=3))
        session.commit()
    with TestClient(app) as c:
        yield c


def names(html: str) -> list[str]:
    found = [reg for reg in ("32600000001", "32600000002", "32600000003") if reg in html]
    return sorted(found, key=html.index)  # в порядке появления на странице


def test_registry_lists_sorted_by_rating(client: TestClient) -> None:
    r = client.get("/")
    assert r.status_code == 200
    html = r.text
    assert html.index("32600000001") < html.index("32600000002")  # 100 раньше 0
    assert "Найдено: <b>3</b>" in html
    assert "dl-urgent" in html  # срок меньше 3 дней подсвечен


def test_registry_filters_and_search(client: TestClient) -> None:
    assert names(client.get("/?law=44").text) == ["32600000002"]
    assert names(client.get("/?q=канцтовар").text) == ["32600000002"]
    assert set(names(client.get("/?min_rating=50").text)) == {"32600000001", "32600000003"}
    assert names(client.get("/?urgent=1").text) == ["32600000001"]
    assert names(client.get("/?sort=price&dir=asc").text)[0] == "32600000002"
    assert names(client.get("/?status=interesting").text) == []
    htmx = client.get("/?law=44", headers={"HX-Request": "true"}).text
    assert "<html" not in htmx and "32600000002" in htmx  # частичная вставка для HTMX


def test_status_change_row_and_card(client: TestClient) -> None:
    with state(client).factory() as s:
        tender = s.scalar(select(Tender).where(Tender.reg_number == "32600000001"))
        assert tender is not None
        tid = tender.id
        pid = s.scalar(select(Profile.id))
    r = client.post(
        f"/tenders/{tid}/status",
        data={"status": "interesting", "profile_id": pid, "view": "row"},
        headers={"HX-Request": "true"},
    )
    assert r.status_code == 200 and r.text.lstrip().startswith("<tr") and "interesting" in r.text
    assert names(client.get("/?status=interesting").text) == ["32600000001"]
    r = client.post(
        f"/tenders/{tid}/status", data={"status": "rejected", "view": "card", "profile_id": pid}
    )
    assert 'id="actions"' in r.text
    assert client.post(f"/tenders/{tid}/status", data={"status": "bogus"}).status_code == 400


def test_card_renders_breakdown(client: TestClient) -> None:
    with state(client).factory() as s:
        tid = s.scalar(select(Tender.id).where(Tender.reg_number == "32600000001"))
    r = client.get(f"/tenders/{tid}")
    assert r.status_code == 200
    for needle in ("Какие правила сработали", "газопровод", "Открыть в ЕИС", "Журнал отправок в CRM",
                   "История изменений", "5 000 000"):  # fmt: skip
        assert needle in r.text
    assert client.get("/tenders/9999").status_code == 404


def test_excel_export_respects_filters(client: TestClient) -> None:
    r = client.get("/export.xlsx?min_rating=50")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(BytesIO(r.content)).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Рейтинг" and len(rows) == 3  # заголовок + 2 закупки
    assert {row[13] for row in rows[1:]} == {"32600000001", "32600000003"}


def test_settings_save_keeps_secrets_blank(client: TestClient) -> None:
    form = {
        "ai_provider": "local", "webhook_url": "http://localhost:5678/webhook/z",
        "webhook_header": "X-Webhook-Secret", "webhook_secret": "abc", "telegram_token": "1:T",
        "telegram_chat_id": "5", "telegram_threshold": "65", "telegram_hour": "8",
        "telegram_enabled": "on",
    }  # fmt: skip
    r = client.post("/settings", data=form, follow_redirects=False)
    assert r.status_code == 303
    with state(client).factory() as s:
        rt = load_runtime(s, state(client).settings)
    assert rt.ai_provider == "local" and rt.webhook_secret == "abc" and rt.telegram_hour == 8
    assert rt.telegram_enabled and rt.telegram_threshold == 65
    # повторное сохранение с пустыми секретами не стирает их
    form.update(webhook_secret="", telegram_token="")
    client.post("/settings", data=form)
    page = client.get("/settings").text
    assert "abc" not in page and "1:T" not in page  # секреты в интерфейс не выводятся
    with state(client).factory() as s:
        rt = load_runtime(s, state(client).settings)
    assert rt.webhook_secret == "abc" and rt.telegram_token == "1:T"
    bad = client.post("/settings", data={**form, "webhook_url": "ftp://x"})
    assert bad.status_code == 400 and "http" in bad.text


def test_demo_footer_marker(client: TestClient) -> None:
    assert "Демо-режим" not in client.get("/").text
    client.post(
        "/settings", data={"ai_provider": "demo", "telegram_threshold": "70", "telegram_hour": "9"}
    )
    assert "Демо-режим" in client.get("/").text


def test_crm_send_with_mocked_webhook(client: TestClient) -> None:
    with state(client).factory() as s:
        tid = s.scalar(select(Tender.id).where(Tender.reg_number == "32600000001"))
        pid = s.scalar(select(Profile.id))
    # без адреса вебхука - понятная ошибка, статус не меняется
    r = client.post(f"/tenders/{tid}/crm", data={"profile_id": pid})
    assert "вебхука" in r.text
    client.post("/settings", data={"ai_provider": "yandex", "webhook_url": "http://n8n.local/hook",
                                   "telegram_threshold": "70", "telegram_hour": "9"})  # fmt: skip
    seen: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request)
        return httpx.Response(
            200, json={"crm": "amocrm", "external_id": "55", "url": "https://c/55"}
        )

    state(client).crm_client_factory = lambda: httpx.Client(transport=httpx.MockTransport(handler))
    r = client.post(f"/tenders/{tid}/crm", data={"profile_id": pid})
    assert r.status_code == 200 and "отправлено" in r.text and "55" in r.text
    assert json.loads(seen[0].content)["reg_number"] == "32600000001"
    with state(client).factory() as s:
        assert s.get(Tender, tid).status == "sent_to_crm"  # type: ignore[union-attr]
        assert len(s.scalars(select(CrmExport)).all()) == 1


def test_profiles_pages_and_rules_editor(client: TestClient) -> None:
    assert "Газ" in client.get("/profiles").text
    with state(client).factory() as s:
        pid = s.scalar(select(Profile.id))
    form = {"name": "Газ 2", "search_string": "газ", "morphology": "on", "fz44": "on", "law44": "on",
            "price_min": "100 000", "stop_words": "аренда, поставка", "min_days": "3", "max_pages": "2",
            "ai_criterion": "проектирование", "ai_weight": "30", "interval_minutes": "60"}  # fmt: skip
    r = client.post(f"/profiles/{pid}", data=form, follow_redirects=False)
    assert r.status_code == 303
    with state(client).factory() as s:
        p = s.get(Profile, pid)
        assert p is not None
        assert p.name == "Газ 2" and p.interval_minutes == 60 and p.ai_weight == 0.3
        assert p.filter["stop_words"] == ["аренда", "поставка"] and p.filter["laws"] == [44]
        assert p.search["fz223"] is False
    # правила: добавить, ошибка веса, удалить
    r = client.post(
        f"/profiles/{pid}/rules", data={"type": "okpd2", "value": "71.12", "weight": "3"}
    )
    assert "71.12" in r.text
    assert (
        "числом"
        in client.post(
            f"/profiles/{pid}/rules", data={"type": "okpd2", "value": "1", "weight": "x"}
        ).text
    )
    with state(client).factory() as s:
        rule_id = s.scalar(select(Rule.id).where(Rule.value == "71.12"))
    client.post(f"/profiles/{pid}/rules/{rule_id}/delete")
    with state(client).factory() as s:
        assert s.scalar(select(Rule.id).where(Rule.value == "71.12")) is None
    # ошибка валидации
    assert client.post("/profiles/new", data={"name": ""}).status_code == 400
    assert client.post(f"/profiles/{pid}/rescore", follow_redirects=False).status_code == 303


def test_profile_import_url(client: TestClient) -> None:
    from .test_search_url import TEST_URL

    r = client.post(
        "/profiles/import", data={"url": TEST_URL, "name": "Импорт"}, follow_redirects=False
    )
    assert r.status_code == 303
    again = client.post("/profiles/import", data={"url": TEST_URL, "name": "Импорт"})
    assert again.status_code == 400 and "уже есть" in again.text


def test_runs_page_and_foreign_hosts_rejected(client: TestClient) -> None:
    assert "Журнал прогонов" in client.get("/runs").text
    assert client.get("/", headers={"host": "evil.example"}).status_code == 400
    r = client.post("/settings", data={}, headers={"origin": "http://evil.example"})
    assert r.status_code == 403


def test_deadline_highlight_rule() -> None:
    today = date(2026, 10, 4)
    assert deadline_info(date(2026, 10, 6), today) == (2, "dl-urgent")
    assert deadline_info(date(2026, 10, 7), today) == (3, "")
    assert deadline_info(date(2026, 10, 3), today) == (-1, "dl-past")
    assert deadline_info(None, today) == (None, "")
