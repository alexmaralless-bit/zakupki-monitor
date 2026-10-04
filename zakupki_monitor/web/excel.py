"""Выгрузка выборки реестра в Excel (openpyxl)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Profile, Score, Tender
from .queries import STATUSES

HEADERS = [
    ("Рейтинг", 10),
    ("По правилам", 12),
    ("ИИ", 8),
    ("Объект закупки", 60),
    ("Заказчик", 40),
    ("ИНН заказчика", 14),
    ("Начальная цена", 18),
    ("Валюта", 8),
    ("Подача до", 12),
    ("Закон", 8),
    ("Профиль", 24),
    ("Статус", 14),
    ("ОКПД2", 20),
    ("Номер", 22),
    ("Ссылка ЕИС", 40),
]


def build_workbook(rows: list[tuple[Tender, Score, Profile]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Закупки"
    ws.append([h for h, _ in HEADERS])
    for i, (_, width) in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2B2F36")
        cell.alignment = Alignment(vertical="center")
    for t, s, p in rows:
        ws.append(
            [
                s.rating,
                s.rules_score,
                s.ai_score,
                t.object,
                t.customer_name,
                t.customer_inn,
                float(t.initial_price) if t.initial_price is not None else None,
                t.currency,
                t.deadline_at,
                f"{t.law}-ФЗ" if t.law else None,
                p.name,
                STATUSES.get(t.status, t.status),
                ", ".join(t.okpd2 or []),
                t.reg_number,
                t.url,
            ]
        )
        r = ws.max_row
        ws.cell(row=r, column=7).number_format = "#,##0.00"
        ws.cell(row=r, column=9).number_format = "DD.MM.YYYY"
        ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        if t.url:
            ws.cell(row=r, column=15).hyperlink = t.url
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
